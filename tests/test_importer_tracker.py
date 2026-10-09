"""Tracker workbooks: several sheets, title rows above the header, recruiter titles, statuses.

Shaped like a real job-search tracker (made-up people and domains; nothing personal here).
"""

import io
import sqlite3

from openpyxl import Workbook

from mailmate import db
from mailmate.importer import analyse, job_sheets
from tests.test_importer import GOOD_ROW, HEADER, csv_bytes


def tracker_bytes() -> bytes:
    """An overview sheet without emails, then contact sheets whose headers sit below title rows,
    recruiter titles in 'Title' / 'Role / Title' / 'Role', and split first/last names."""
    wb = Workbook()
    overview = wb.active
    overview.title = "Company Tracker"
    for r in (["Company Tracker"], [], ["#", "Company", "Tier", "Contact Emails"], ["1", "Nimbus Labs", "Tier 1", ""]):
        overview.append(r)
    hunter = wb.create_sheet("Hunter Contacts")
    for r in (["Hunter Contacts: recruiter emails"], ["status 'valid' means..."], [],
              ["#", "Company", "Name", "Title", "Email", "Hunter status", "Note"],
              ["1", "Nimbus Labs", "Priya Sharma", "Senior Technical Recruiter", "priya@nimbuslabs.ai", "valid", ""],
              ["2", "Quillstack", "Neha Gupta", "Talent Partner", "neha@quillstack.io", "accept_all", ""],
              ["3", "Quillstack", "Old Lead", "Recruiter", "gone@quillstack.io", "invalid", ""],
              ["4", "Nimbus Labs", "Ravi Kumar", "Head of Talent", "ravi@nimbuslabs.ai", "valid", ""],
              [None, None, "A note typed below the table", None, None, None, None]):
        hunter.append(r)
    outreach = wb.create_sheet("Outreach Tracker")
    for r in (["Outreach Tracker"], [],
              ["#", "First Name", "Last Name", "Company", "Role / Title", "Email", "Status", "Date Sent"],
              ["Ex.", "Sample", "Person", "Example AI Co", "Technical Recruiter", "sample@example.com", "Sent",
               "2026-10-01"],
              ["1", "Priya", "Sharma", "Nimbus Labs", "Senior Technical Recruiter", "PRIYA@nimbuslabs.ai",
               "To Contact", ""],
              ["2", "Ravi", "Kumar", "Nimbus Labs", "Head of Talent", "ravi@nimbuslabs.ai", "Sent", "2026-10-02"],
              ["3", "Arjun", "Rao", "Quillstack", "Hiring @ Quillstack (recruiter)", "arjun@quillstack.io",
               "To Contact", ""]):
        outreach.append(r)
    found = wb.create_sheet("Emails Found")
    for r in ([], ["#", "Company", "Name", "Role", "Email", "Status"],
              ["1", "Quillstack", "—", "HR (published for resumes)", "hr@quillstack.io", "Published by the company"],
              ["2", "Nimbus Labs", "—", "General contact", "hello@nimbuslabs.ai", "Published by the company"],
              ["", "Hugging Face", "", "", "", "no published address"]):
        found.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def tracker(conn, fake_dns, **kw):
    report = analyse("tracker.xlsx", tracker_bytes(), conn, fake_dns, **{"default_role": "AI Engineer", **kw})
    assert not report.file_errors, report.file_errors
    return report, {(r.sheet, r.row): r for r in report.rows}


def test_job_sheets_are_the_ones_with_company_and_email():
    assert job_sheets("tracker.xlsx", tracker_bytes()) == ["Hunter Contacts", "Outreach Tracker", "Emails Found"]
    assert job_sheets("jobs.csv", csv_bytes(HEADER, GOOD_ROW)) == []


def test_headers_below_title_rows_are_found(conn, fake_dns):
    report, _ = tracker(conn, fake_dns)
    assert [(s.name, s.header_row) for s in report.sheets] == [
        ("Hunter Contacts", 4), ("Outreach Tracker", 3), ("Emails Found", 2)]


def test_recruiter_titles_never_become_the_role(conn, fake_dns):
    report, rows = tracker(conn, fake_dns)
    priya = rows[("Hunter Contacts", 5)].contact
    assert (priya.role, priya.title, priya.name) == ("AI Engineer", "Senior Technical Recruiter", "Priya Sharma")
    hr = rows[("Emails Found", 3)].contact
    assert (hr.role, hr.title, hr.name) == ("AI Engineer", "HR (published for resumes)", "")
    found = next(s for s in report.sheets if s.name == "Emails Found")
    assert found.columns["title"] == "Role" and "role" not in found.columns
    assert "mostly holds recruiter/HR titles" in found.notes[0]


def test_a_target_role_column_wins_over_the_default(conn, fake_dns):
    data = csv_bytes(["Company", "Target Role", "Title", "Email"],
                     ["Nimbus Labs", "Agentic AI Engineer", "Recruiter", "priya@nimbuslabs.ai"],
                     ["Quillstack", "", "Recruiter", "neha@quillstack.io"])
    report = analyse("jobs.csv", data, conn, fake_dns, default_role="AI Engineer")
    assert [r.contact.role for r in report.valid] == ["Agentic AI Engineer", "AI Engineer"]


def test_without_any_role_the_row_says_what_to_do(conn, fake_dns):
    report = analyse("jobs.csv", csv_bytes(["Company", "Email"], ["Nimbus Labs", "priya@nimbuslabs.ai"]), conn, fake_dns)
    assert "type the role you're applying for" in report.rejected[0].errors[0]


def test_first_and_last_name_columns_make_the_name(conn, fake_dns):
    _, rows = tracker(conn, fake_dns, sheets=["Outreach Tracker"])
    assert rows[("Outreach Tracker", 7)].contact.name == "Arjun Rao"


def test_email_checker_status(conn, fake_dns):
    _, rows = tracker(conn, fake_dns)
    neha = rows[("Hunter Contacts", 6)]
    assert neha.ok and "may bounce" in neha.warnings[0]
    assert rows[("Hunter Contacts", 7)].errors == ["the email checker says this address is invalid"]


def test_already_contacted_is_refused_in_every_sheet(conn, fake_dns):
    report, rows = tracker(conn, fake_dns)
    assert rows[("Outreach Tracker", 6)].errors[0].startswith("already contacted outside MailMate (Status 'Sent')")
    assert rows[("Hunter Contacts", 8)].errors == ["already contacted outside MailMate (another row's Status says so)"]
    assert report.contacted_elsewhere == ["ravi@nimbuslabs.ai"]


def test_sample_rows_with_example_addresses_are_refused(conn, fake_dns):
    report, rows = tracker(conn, fake_dns)
    assert "example address" in rows[("Outreach Tracker", 4)].errors[0]
    assert "sample@example.com" not in report.contacted_elsewhere


def test_duplicates_across_sheets_name_the_sheet(conn, fake_dns):
    _, rows = tracker(conn, fake_dns)
    assert rows[("Outreach Tracker", 5)].errors == ["duplicate of 'Hunter Contacts' row 5"]


def test_note_rows_are_skipped_and_rows_missing_an_email_are_reported(conn, fake_dns):
    report, rows = tracker(conn, fake_dns)
    assert next(s for s in report.sheets if s.name == "Hunter Contacts").blank_rows == 1
    assert rows[("Emails Found", 5)].errors == ["email is missing"]


def test_whole_tracker_summary(conn, fake_dns):
    report, _ = tracker(conn, fake_dns)
    assert sorted(r.contact.email for r in report.valid) == [
        "arjun@quillstack.io", "hello@nimbuslabs.ai", "hr@quillstack.io", "neha@quillstack.io", "priya@nimbuslabs.ai"]


def test_only_the_ticked_sheets_are_read(conn, fake_dns):
    report, _ = tracker(conn, fake_dns, sheets=["Emails Found"])
    assert {r.sheet for r in report.rows} == {"Emails Found"}
    assert analyse("tracker.xlsx", tracker_bytes(), conn, fake_dns, sheets=[]).file_errors == ["no sheet selected"]


def test_title_is_stored_and_old_databases_get_the_column(tmp_path, fake_dns):
    path = tmp_path / "old.db"
    old = sqlite3.connect(path)
    old.execute("CREATE TABLE contacts (id INTEGER PRIMARY KEY, company TEXT NOT NULL, role TEXT NOT NULL, "
                "requirements TEXT NOT NULL DEFAULT '', email TEXT NOT NULL UNIQUE COLLATE NOCASE, "
                "name TEXT NOT NULL DEFAULT '', job_url TEXT NOT NULL DEFAULT '', status TEXT NOT NULL DEFAULT 'new', "
                "source_file TEXT NOT NULL DEFAULT '', created_at TEXT NOT NULL)")
    old.close()
    conn = db.connect(path)
    report, _ = tracker(conn, fake_dns, sheets=["Hunter Contacts"])
    db.add_contacts(conn, [r.contact for r in report.valid], "tracker.xlsx")
    title = conn.execute("SELECT title FROM contacts WHERE email = 'priya@nimbuslabs.ai'").fetchone()[0]
    conn.close()
    assert title == "Senior Technical Recruiter"
