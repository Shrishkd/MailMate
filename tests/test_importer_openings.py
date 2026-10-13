"""A tracker with a sheet of openings (links) apart from the contact sheets, 'Apply here' hyperlinks,
sheets named 'Applied', a Group column, and email ratings in words. Made-up companies and people."""

import io

import pytest
from openpyxl import Workbook

from mailmate import db
from mailmate.importer import Contact, analyse, job_sheets, pick_link, JobLink

NIMBUS_JOB = "https://jobs.example.org/nimbus/42"
NIMBUS_CHAT = "https://jobs.example.org/nimbus/chat"
QUILL_JOB = "https://jobs.example.org/quillstack/7"


def workbook(*, openings_status=("Not applied", "Not applied", "Applied")) -> bytes:
    wb = Workbook()
    openings = wb.active
    openings.title = "Openings"
    openings.append(["New openings"])
    openings.append([])
    openings.append(["#", "Company", "Role", "Direct apply link", "Status", "Applied on"])
    for i, (company, role, url, status) in enumerate([
            ("Nimbus Labs", "GenAI Engineer Intern", NIMBUS_JOB, openings_status[0]),
            ("Nimbus Labs", "AI Chatbot Developer Intern", NIMBUS_CHAT, openings_status[1]),
            ("Quillstack", "AI Engineer", QUILL_JOB, openings_status[2])], 1):
        openings.append([i, company, role, "Apply here", status, None])
        openings.cell(row=3 + i, column=4).hyperlink = url            # the text says 'Apply here'
    openings.append([None, "A note below the table", None, None, None, None])

    new = wb.create_sheet("Opening Contacts")
    new.append(["Contacts"])
    new.append(["#", "Company", "Contact name", "Role", "Email", "How the email was found"])
    new.append([1, "Nimbus Labs", "Priya Sharma", "AI Chatbot Developer Intern", "priya@nimbuslabs.ai",
                "Hunter pattern-generated, verified valid"])
    new.append([2, "Nimbus Labs", "Ravi Kumar", "GenAI Engineer Intern", "ravi@nimbuslabs.ai",
                "Hunter pattern-generated, not verified"])
    new.append([3, "Quillstack", "Neha Gupta", "AI Engineer", "neha@quillstack.io",
                "Hunter pattern-generated, accept-all domain (cannot confirm)"])
    new.append([4, "Quillstack", "A Founder", "AI Engineer", None, "No email found"])

    done = wb.create_sheet("Applied-5 Contacts")
    done.append(["Contacts for companies I applied to"])
    done.append(["#", "Company", "Contact name", "Role", "Email", "How the email was found"])
    done.append([1, "Zephyr Works", "Anita Rao", "Agentic AI Intern", "anita@quillstack.io", "Published on company page"])

    tracker = wb.create_sheet("Outreach Tracker")
    tracker.append(["Outreach"])
    tracker.append(["#", "Group", "Company", "Contact", "Email", "Email basis", "Status"])
    tracker.append([1, "Applied", "Brightfield", "Vikram Nair", "vikram@nimbuslabs.ai", "Published", "To Contact"])
    tracker.append([2, "New opening", "Nimbus Labs", "Priya Sharma", "priya@nimbuslabs.ai", "Hunter, valid", "To Contact"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def report(conn, fake_dns, **kw):
    r = analyse("tracker.xlsx", workbook(**kw), conn, fake_dns, default_role="AI Engineer")
    assert not r.file_errors, r.file_errors
    return r


def contacts(r) -> dict[str, Contact]:
    return {x.contact.email: x.contact for x in r.valid}


# --- hyperlinks and sheets of openings ----------------------------------------------------------------

def test_the_address_behind_apply_here_is_read(conn, fake_dns):
    r = report(conn, fake_dns)
    assert sorted(l.url for l in r.job_links) == sorted([NIMBUS_JOB, NIMBUS_CHAT, QUILL_JOB])
    assert "Apply here" not in str([l.url for l in r.job_links])


def test_a_sheet_of_openings_imports_nobody_and_says_what_it_found(conn, fake_dns):
    r = report(conn, fake_dns)
    assert job_sheets("tracker.xlsx", workbook()) == ["Opening Contacts", "Applied-5 Contacts", "Outreach Tracker"]
    note = next(s for s in r.sheets if s.name == "Openings").notes[0]
    assert note.startswith("3 job link(s) found (1 marked applied)") and "nobody is imported" in note
    assert {x.sheet for x in r.rows} == {"Opening Contacts", "Applied-5 Contacts", "Outreach Tracker"}


def test_contacts_get_the_link_of_their_opening(conn, fake_dns):
    got = contacts(report(conn, fake_dns))
    assert got["priya@nimbuslabs.ai"].job_url == NIMBUS_CHAT        # her role names the chatbot opening
    assert got["ravi@nimbuslabs.ai"].job_url == NIMBUS_JOB          # his role names the GenAI opening
    assert got["neha@quillstack.io"].job_url == QUILL_JOB
    assert got["anita@quillstack.io"].job_url == ""                 # Zephyr Works has no opening listed


def test_the_link_is_the_first_of_the_company_when_no_role_matches():
    links = [JobLink(company="A", role="Data Intern", url="u1"), JobLink(company="A", role="ML Intern", url="u2")]
    assert pick_link(links, "Research Scientist").url == "u1"
    assert pick_link(links, "ML Intern / Data Intern").url == "u1"
    assert pick_link([], "x") is None


def test_link_sheets_are_read_even_when_only_one_sheet_is_ticked(conn, fake_dns):
    r = analyse("tracker.xlsx", workbook(), conn, fake_dns, sheets=["Opening Contacts"], default_role="AI Engineer")
    assert len(r.job_links) == 3 and {x.sheet for x in r.rows} == {"Opening Contacts"}


# --- applied by direct link ----------------------------------------------------------------------------

def test_a_sheet_named_applied_marks_its_companies(conn, fake_dns):
    r = report(conn, fake_dns)
    assert contacts(r)["anita@quillstack.io"].applied is True
    sheet = next(s for s in r.sheets if s.name == "Applied-5 Contacts")
    assert "so every company in it is marked as applied by direct link" in sheet.notes[0]


def test_an_applied_status_on_an_opening_marks_that_company_and_its_contacts(conn, fake_dns):
    r = report(conn, fake_dns)
    got = contacts(r)
    assert got["neha@quillstack.io"].applied is True                  # Quillstack's opening says Applied
    assert got["priya@nimbuslabs.ai"].applied is False                # Nimbus: "Not applied" twice
    assert "Quillstack" in r.applied_companies and "Nimbus Labs" not in r.applied_companies


def test_the_group_column_marks_applied(conn, fake_dns):
    r = report(conn, fake_dns)
    assert contacts(r)["vikram@nimbuslabs.ai"].applied is True
    assert "Brightfield" in r.applied_companies
    assert contacts(r)["priya@nimbuslabs.ai"].applied is False        # 'New opening' is not applied


def test_not_applied_never_counts_as_applied(conn, fake_dns):
    r = report(conn, fake_dns, openings_status=("Not applied",) * 3)
    assert contacts(r)["neha@quillstack.io"].applied is False
    assert sorted(r.applied_companies) == ["Brightfield", "Zephyr Works"]


# --- the rest of this layout ------------------------------------------------------------------------------

def test_email_ratings_in_words_give_the_may_bounce_warning(conn, fake_dns):
    r = report(conn, fake_dns)
    by_email = {x.contact.email: x for x in r.valid}
    assert by_email["priya@nimbuslabs.ai"].warnings == []                        # "verified valid"
    assert "may bounce" in by_email["ravi@nimbuslabs.ai"].warnings[0]            # "not verified"
    assert "may bounce" in by_email["neha@quillstack.io"].warnings[0]            # "accept-all ... cannot confirm"
    assert by_email["anita@quillstack.io"].warnings == []                        # "Published"


def test_rows_without_an_email_are_rejected_for_that_reason_only(conn, fake_dns):
    r = report(conn, fake_dns)
    assert [x.errors for x in r.rejected if "A Founder" in str(x.raw)] == [["email is missing"]]


def test_the_roles_come_from_the_sheet_not_from_the_default(conn, fake_dns):
    assert contacts(report(conn, fake_dns))["anita@quillstack.io"].role == "Agentic AI Intern"


def test_existing_contacts_get_the_link_from_the_openings_sheet(conn, fake_dns):
    db.add_contacts(conn, [Contact(company="Quillstack", role="AI Engineer", email="neha@quillstack.io")], "old.csv")
    r = report(conn, fake_dns)
    assert {"email": "neha@quillstack.io", "job_url": QUILL_JOB, "priority": ""} in r.known_updates
    assert db.fill_contact_blanks(conn, r.known_updates, "tracker.xlsx") == 1
    assert conn.execute("SELECT job_url FROM contacts WHERE email = 'neha@quillstack.io'").fetchone()[0] == QUILL_JOB


@pytest.mark.parametrize("name", ["Applied", "Applied-20 Contacts", "Companies I applied to"])
def test_sheet_names_that_say_applied(name):
    from mailmate.importer import _sheet_says_applied
    assert _sheet_says_applied(name)


@pytest.mark.parametrize("name", ["Not applied", "Openings", "Unapplied roles", ""])
def test_sheet_names_that_do_not(name):
    from mailmate.importer import _sheet_says_applied
    assert not _sheet_says_applied(name)
