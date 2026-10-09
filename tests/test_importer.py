import csv
import io
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from mailmate import db
from mailmate.importer import analyse

FIXTURES = Path(__file__).parent / "fixtures"

HEADER = ["Company", "Role", "Requirements", "Recruiter Email", "Recruiter Name", "Job URL"]
GOOD_ROW = ["Nimbus Labs", "GenAI Engineer", "Python, RAG", "priya@nimbuslabs.ai", "Priya", "https://nimbuslabs.ai/j/1"]


def csv_bytes(*rows, sep=",", encoding="utf-8") -> bytes:
    buf = io.StringIO()
    csv.writer(buf, delimiter=sep, lineterminator="\n").writerows(rows)
    return buf.getvalue().encode(encoding)


def xlsx_bytes(*rows) -> bytes:
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def docx_bytes(*rows, intro_table=None) -> bytes:
    doc = Document()
    doc.add_paragraph("Jobs I'm applying to")
    for table_rows in ([intro_table] if intro_table else []) + [list(rows)]:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for i, r in enumerate(table_rows):
            for j, v in enumerate(r):
                table.cell(i, j).text = v
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def one_row(conn, fake_dns, **overrides):
    """Analyse a CSV with a single job row; overrides replace fields of GOOD_ROW by header."""
    row = list(GOOD_ROW)
    keys = ["company", "role", "requirements", "email", "name", "job_url"]
    for k, v in overrides.items():
        row[keys.index(k)] = v
    report = analyse("jobs.csv", csv_bytes(HEADER, row), conn, fake_dns)
    assert not report.file_errors, report.file_errors
    return report.rows[0]


# --- file types -----------------------------------------------------------------------------

def test_csv(conn, fake_dns):
    report = analyse("jobs.csv", csv_bytes(HEADER, GOOD_ROW), conn, fake_dns)
    [row] = report.valid
    assert row.contact.model_dump() == {
        "company": "Nimbus Labs", "role": "GenAI Engineer", "requirements": "Python, RAG",
        "email": "priya@nimbuslabs.ai", "name": "Priya", "title": "", "job_url": "https://nimbuslabs.ai/j/1"}


def test_xlsx(conn, fake_dns):
    report = analyse("jobs.xlsx", xlsx_bytes(HEADER, GOOD_ROW), conn, fake_dns)
    assert [r.contact.email for r in report.valid] == ["priya@nimbuslabs.ai"]


def test_docx_uses_the_table_with_job_headers(conn, fake_dns):
    data = docx_bytes(HEADER, GOOD_ROW, intro_table=[["Week", "Goal"], ["1", "Apply to 20 jobs"]])
    report = analyse("jobs.docx", data, conn, fake_dns)
    assert [r.contact.company for r in report.valid] == ["Nimbus Labs"]


def test_csv_saved_by_excel_in_cp1252(conn, fake_dns):
    data = csv_bytes(HEADER, ["Société Générale"] + GOOD_ROW[1:], encoding="cp1252")
    assert analyse("jobs.csv", data, conn, fake_dns).valid[0].contact.company == "Société Générale"


def test_semicolon_separated_csv(conn, fake_dns):
    report = analyse("jobs.csv", csv_bytes(HEADER, GOOD_ROW, sep=";"), conn, fake_dns)
    assert len(report.valid) == 1


def test_unsupported_file_type(conn, fake_dns):
    report = analyse("jobs.pdf", b"%PDF-1.4", conn, fake_dns)
    assert report.file_errors == ["unsupported file type '.pdf'; use .csv, .xlsx, .docx"]


def test_corrupt_xlsx(conn, fake_dns):
    report = analyse("jobs.xlsx", b"this is not a workbook", conn, fake_dns)
    assert "couldn't open the file" in report.file_errors[0]


def test_empty_file(conn, fake_dns):
    assert analyse("jobs.csv", b"", conn, fake_dns).file_errors == ["the file is empty"]


def test_header_without_jobs(conn, fake_dns):
    assert analyse("jobs.csv", csv_bytes(HEADER), conn, fake_dns).file_errors == ["the file has a header but no jobs"]


# --- columns --------------------------------------------------------------------------------

def test_missing_required_column(conn, fake_dns):
    report = analyse("jobs.csv", csv_bytes(["Company", "Role", "Notes"], ["A", "B", "C"]), conn, fake_dns)
    assert len(report.file_errors) == 1
    assert report.file_errors[0].startswith("no 'Recruiter email' column found")


def test_two_columns_for_the_same_field(conn, fake_dns):
    report = analyse("jobs.csv", csv_bytes(["Company", "Role", "Email", "Recruiter Email"], ["A", "B", "a@b.com", "c@d.com"]),
                     conn, fake_dns)
    assert report.file_errors == ["two columns mean 'Recruiter email': 'Email' and 'Recruiter Email'; rename or delete one"]


def test_unknown_columns_are_listed_not_fatal(conn, fake_dns):
    report = analyse("jobs.csv", csv_bytes(HEADER + ["Notes"], GOOD_ROW + ["x"]), conn, fake_dns)
    assert report.ignored_columns == ["Notes"] and len(report.valid) == 1


# --- row checks -----------------------------------------------------------------------------

@pytest.mark.parametrize("field, value, error", [
    ("company", "", "company is missing"),
    ("company", "N/A", "company is missing"),
    ("role", "-", "role is missing"),
    ("email", "", "email is missing"),
    ("email", "not-an-email", "not a valid email address"),
    ("email", "a@nimbuslabs.ai, b@nimbuslabs.ai", "more than one address in the cell"),
    ("email", "no-reply@nimbuslabs.ai", "no-reply address"),
    ("email", "jobs@deadcorp.com", "can't receive email: the domain does not exist"),
    ("email", "jobs@parkedsite.com", "can't receive email: the domain says it accepts no email (null MX)"),
    ("company", "x" * 151, "company is 151 characters long"),
])
def test_row_errors(conn, fake_dns, field, value, error):
    row = one_row(conn, fake_dns, **{field: value})
    assert not row.ok and row.contact is None
    assert any(e.startswith(error) for e in row.errors), row.errors


@pytest.mark.parametrize("field, value, warning", [
    ("email", "careers@slowdns.in", "couldn't check the domain (Timeout); it will be checked again before sending"),
    ("job_url", "see LinkedIn", "job URL 'see LinkedIn' doesn't look like a web link"),
])
def test_row_warnings_still_import(conn, fake_dns, field, value, warning):
    row = one_row(conn, fake_dns, **{field: value})
    assert row.ok and row.warnings == [warning]


def test_email_is_cleaned(conn, fake_dns):
    assert one_row(conn, fake_dns, email="  mailto:Priya@NimbusLabs.AI ").contact.email == "priya@nimbuslabs.ai"
    named = one_row(conn, fake_dns, email="Priya Sharma <priya@nimbuslabs.ai>", name="")
    assert (named.contact.email, named.contact.name) == ("priya@nimbuslabs.ai", "Priya Sharma")


def test_duplicate_in_file_ignores_case(conn, fake_dns):
    data = csv_bytes(HEADER, GOOD_ROW, GOOD_ROW[:3] + ["PRIYA@nimbuslabs.ai"] + GOOD_ROW[4:])
    report = analyse("jobs.csv", data, conn, fake_dns)
    assert [r.row for r in report.valid] == [2]
    assert report.rejected[0].errors == ["duplicate of row 2"]


def test_already_a_contact(conn, fake_dns):
    first = analyse("jobs.csv", csv_bytes(HEADER, GOOD_ROW), conn, fake_dns)
    assert db.add_contacts(conn, [r.contact for r in first.valid], "jobs.csv") == (1, [])
    row = one_row(conn, fake_dns, email="Priya@NimbusLabs.ai")
    assert row.errors[0].startswith("already a contact (added ")
    assert "for Nimbus Labs); MailMate never emails an address twice" in row.errors[0]


def test_do_not_contact(conn, fake_dns):
    db.suppress(conn, "Priya@nimbuslabs.ai", "not_interested")
    assert one_row(conn, fake_dns).errors == ["on the do-not-contact list (not interested)"]


def test_blank_rows_are_skipped_and_row_numbers_match_excel(conn, fake_dns):
    data = csv_bytes(HEADER, ["", "", "", "", "", ""], GOOD_ROW)
    report = analyse("jobs.csv", data, conn, fake_dns)
    assert report.blank_rows == 1 and report.valid[0].row == 3


# --- the messy sample file --------------------------------------------------------------------

def test_messy_file_gives_a_readable_report(conn, fake_dns):
    db.suppress(conn, "stop@quillstack.io", "bounce")
    db.add_contacts(conn, [one_row(conn, fake_dns, email="existing@nimbuslabs.ai").contact], "earlier.csv")

    report = analyse("messy_jobs.csv", (FIXTURES / "messy_jobs.csv").read_bytes(), conn, fake_dns)

    assert report.ignored_columns == ["Notes"]
    assert report.blank_rows == 1
    by_row = {r.row: r for r in report.rows}
    assert [r.row for r in report.valid] == [3, 4, 9, 16, 17]
    assert by_row[4].contact.email == "hr@quillstack.io" and by_row[4].contact.role == "AI Engineer"
    expected_errors = {
        5: "company is missing",
        6: "can't receive email: the domain does not exist",
        7: "can't receive email: the domain says it accepts no email (null MX)",
        8: "duplicate of row 3",
        10: "no-reply address; nobody reads it",
        11: "role is missing",
        12: "not a valid email address",
        14: "on the do-not-contact list (bounce)",
        15: "already a contact",
    }
    for row, error in expected_errors.items():
        assert by_row[row].errors[0].startswith(error), (row, by_row[row].errors)
    assert "more than one address in the cell" in by_row[11].errors[1]
    assert "couldn't check the domain" in by_row[9].warnings[0]

    text = report.as_text()
    assert "5 will be imported, 9 rejected, 1 blank or note rows skipped" in text
    assert "Recruiter email <- 'Recruiter E-mail'" in text
